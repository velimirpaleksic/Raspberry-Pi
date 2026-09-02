from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document

from project.core import config
from project.core.runtime_settings import get_selected_printer
from project.services.storage_cleanup import DiskInfo, collect_storage_report, format_bytes
from project.utils.docs.docx_replace_placeholders import replace_dynamic_text
from project.utils.docs.pdf_converter import convert_docx_to_pdf
from project.utils.logging_utils import log_error, log_info
from project.utils.network_status import collect_network_diagnostics
from project.utils.printing.printer_status import collect_printer_diagnostics


REQUIRED_PLACEHOLDERS = (
    "{{DANASNJI_DATUM}}",
    "{{IME}}",
    "{{RODITELJ}}",
    "{{DATUM_RODJENJA}}",
    "{{MJESTO}}",
    "{{OPSTINA}}",
    "{{RAZRED}}",
    "{{STRUKA}}",
    "{{RAZLOG}}",
)


@dataclass(frozen=True)
class SelfTestCheck:
    name: str
    ok: bool
    detail: str


@dataclass(frozen=True)
class SelfTestReport:
    checks: tuple[SelfTestCheck, ...]
    elapsed_seconds: float
    paper_printed: bool = False

    @property
    def ok(self) -> bool:
        return bool(self.checks) and all(check.ok for check in self.checks)


def _compact(value: object, limit: int = 350) -> str:
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 3] + "..."


def format_self_test_report(report: SelfTestReport) -> str:
    title = "✅ Uvjerenja Terminal self-test: USPJEŠAN" if report.ok else "❌ Uvjerenja Terminal self-test: GREŠKA"
    lines = [
        title,
        f"Vrijeme: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}",
        f"Trajanje: {report.elapsed_seconds:.1f}s",
        "Papir: NIJE poslan na štampu",
        "",
    ]
    for check in report.checks:
        marker = "✅" if check.ok else "❌"
        lines.append(f"{marker} {check.name}: {_compact(check.detail)}")
    return "\n".join(lines)


def _iter_document_text(document: Document):
    for paragraph in document.paragraphs:
        yield paragraph.text
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph.text


def _validate_zip(path: Path) -> None:
    if not path.is_file() or path.stat().st_size <= 0:
        raise FileNotFoundError(path)
    with ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise BadZipFile(f"damaged member: {bad_member}")
        if "[Content_Types].xml" not in archive.namelist():
            raise BadZipFile("missing [Content_Types].xml")


def _validate_template(path: Path) -> str:
    _validate_zip(path)
    document = Document(path)
    text = "\n".join(_iter_document_text(document))
    missing = [placeholder for placeholder in REQUIRED_PLACEHOLDERS if placeholder not in text]
    if missing:
        raise ValueError("nedostaju placeholderi: " + ", ".join(missing))
    if text.count("2026/2027") != 1 or "2025/2026" in text:
        raise ValueError("školska godina u templateu nije tačno 2026/2027")
    return f"validan DOCX/OOXML, {len(REQUIRED_PLACEHOLDERS)} obaveznih polja, školska 2026/2027"


def _sample_placeholders() -> dict[str, str]:
    return {
        "{{DANASNJI_DATUM}}": datetime.now().strftime("%d.%m.%Y"),
        "{{IME}}": "АЛЕКСАНДАР МАКСИМИЛИЈАН ПЕТРОВИЋ",
        "{{IME_PREZIME}}": "АЛЕКСАНДАР МАКСИМИЛИЈАН ПЕТРОВИЋ",
        "{{IME_UCENIKA}}": "АЛЕКСАНДАР",
        "{{PREZIME}}": "МАКСИМИЛИЈАН ПЕТРОВИЋ",
        "{{RODITELJ}}": "АЛЕКСАНДРИЈА",
        "{{DATUM_RODJENJA}}": "31.12.2008",
        "{{MJESTO}}": "ДОЊЕ НОВО СЕЛО",
        "{{OPSTINA}}": "ИСТОЧНА ИЛИЏА",
        "{{RAZRED}}": "ЧЕТВРТИ",
        "{{STRUKA}}": max(config.STRUKE, key=len).upper(),
        "{{RAZLOG}}": max(config.RAZLOZI, key=len).upper(),
    }


def _validate_generated_docx(path: Path) -> str:
    _validate_zip(path)
    document = Document(path)
    text = "\n".join(_iter_document_text(document))
    unresolved = sorted(set(re.findall(r"\{\{[^{}]+\}\}", text)))
    if unresolved:
        raise ValueError("neriješeni placeholderi: " + ", ".join(unresolved))
    if "2026/2027" not in text:
        raise ValueError("školska godina nedostaje u generisanom dokumentu")
    return f"generisan i ponovo otvoren, bez placeholdera ({format_bytes(path.stat().st_size)})"


def _pdf_page_count(path: Path, pdfinfo_command: str) -> int:
    completed = subprocess.run(
        [pdfinfo_command, str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        timeout=config.SUBPROCESS_TIMEOUT,
    )
    if completed.returncode != 0:
        detail = completed.stderr or completed.stdout or f"exit {completed.returncode}"
        raise RuntimeError("pdfinfo nije uspio: " + _compact(detail))
    match = re.search(r"^Pages:\s*(\d+)\s*$", completed.stdout or "", flags=re.MULTILINE | re.IGNORECASE)
    if not match:
        raise RuntimeError("pdfinfo nije vratio broj stranica")
    return int(match.group(1))


def _validate_pdf(path: Path, pdfinfo_command: str) -> str:
    if not path.is_file() or path.stat().st_size <= 0:
        raise FileNotFoundError(path)
    data = path.read_bytes()
    if not data.startswith(b"%PDF-") or b"%%EOF" not in data[-2048:]:
        raise ValueError("generisani PDF nema ispravnu PDF strukturu")
    page_count = _pdf_page_count(path, pdfinfo_command)
    if page_count != 1:
        raise ValueError(f"očekivana je 1 stranica, dobijeno {page_count}")
    return f"validan PDF, 1 stranica ({format_bytes(path.stat().st_size)})"


def _storage_check(report: dict[str, DiskInfo]) -> tuple[bool, str]:
    problems: list[str] = []
    summaries: list[str] = []
    warning_percent = max(1, config.STORAGE_ALERT_USED_PERCENT)
    min_free = max(0, config.STORAGE_ALERT_MIN_FREE_MB) * 1024 * 1024

    for name, info in report.items():
        label = "/" if name == "root" else "app data"
        if info.error:
            problems.append(f"{label}: {info.error}")
            continue
        summaries.append(f"{label} {format_bytes(info.free)} slobodno ({info.used_percent:.0f}% zauzeto)")
        if info.used_percent >= warning_percent:
            problems.append(f"{label} je popunjen {info.used_percent:.0f}%")
        if min_free and info.free <= min_free:
            problems.append(f"{label} ima samo {format_bytes(info.free)} slobodno")

    if problems:
        return False, "; ".join(problems)
    return True, "; ".join(summaries) or "prostor nije moguće provjeriti"


def _printer_check() -> tuple[bool, str]:
    info = collect_printer_diagnostics(get_selected_printer())
    resolved = str(info.get("resolved") or "").strip()
    if info.get("ready"):
        return True, f"{resolved or 'printer'} je spreman; testni posao nije poslan"
    code = str(info.get("ready_code") or info.get("detect_code") or "PRN_CHECK_FAILED")
    detail = str(info.get("ready_message") or info.get("detect_message") or "printer nije spreman")
    return False, f"{code}: {detail}"


def _network_check() -> tuple[bool, str]:
    info = collect_network_diagnostics()
    ssid = str(info.get("ssid") or "nepoznat SSID")
    ip = str(info.get("ip") or "bez IP adrese")
    detail = str(info.get("internet_message") or "")
    if info.get("internet"):
        return True, f"internet dostupan; {ssid}; {ip}"
    return False, f"internet nije dostupan; {ssid}; {ip}; {detail}"


def run_self_test() -> SelfTestReport:
    """Check the complete local document path and dependencies without printing."""

    started = time.monotonic()
    checks: list[SelfTestCheck] = []
    template_ok = False
    generated_docx: Path | None = None

    try:
        detail = _validate_template(config.TEMPLATE_FILE)
        template_ok = True
        checks.append(SelfTestCheck("Template", True, detail))
    except Exception as exc:
        checks.append(SelfTestCheck("Template", False, _compact(exc)))

    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    pdfinfo = shutil.which("pdfinfo")
    dependency_problems = []
    if not libreoffice:
        dependency_problems.append("LibreOffice nije pronađen")
    if not pdfinfo:
        dependency_problems.append("pdfinfo/poppler-utils nije pronađen")
    checks.append(
        SelfTestCheck(
            "Programi",
            not dependency_problems,
            "; ".join(dependency_problems) if dependency_problems else "LibreOffice i pdfinfo su dostupni",
        )
    )

    try:
        with tempfile.TemporaryDirectory(prefix="selftest-", dir=str(config.VAR_DIR)) as temp_dir:
            temp_path = Path(temp_dir)
            if not template_ok:
                checks.append(SelfTestCheck("DOCX generisanje", False, "preskočeno zbog neispravnog templatea"))
                checks.append(SelfTestCheck("PDF generisanje", False, "preskočeno zbog neispravnog templatea"))
            else:
                generated_docx = temp_path / "selftest.docx"
                try:
                    replace_dynamic_text(str(config.TEMPLATE_FILE), str(generated_docx), _sample_placeholders())
                    checks.append(SelfTestCheck("DOCX generisanje", True, _validate_generated_docx(generated_docx)))
                except Exception as exc:
                    generated_docx = None
                    checks.append(SelfTestCheck("DOCX generisanje", False, _compact(exc)))

                if generated_docx is None:
                    checks.append(SelfTestCheck("PDF generisanje", False, "preskočeno jer DOCX nije generisan"))
                elif not libreoffice or not pdfinfo:
                    checks.append(SelfTestCheck("PDF generisanje", False, "preskočeno jer nedostaje LibreOffice ili pdfinfo"))
                else:
                    try:
                        pdf_path = Path(convert_docx_to_pdf(str(generated_docx), output_dir=str(temp_path)))
                        checks.append(SelfTestCheck("PDF generisanje", True, _validate_pdf(pdf_path, pdfinfo)))
                    except Exception as exc:
                        checks.append(SelfTestCheck("PDF generisanje", False, _compact(exc)))
    except Exception as exc:
        checks.append(SelfTestCheck("Privremeni direktorij", False, _compact(exc)))

    try:
        printer_ok, printer_detail = _printer_check()
        checks.append(SelfTestCheck("Printer/CUPS", printer_ok, _compact(printer_detail)))
    except Exception as exc:
        checks.append(SelfTestCheck("Printer/CUPS", False, _compact(exc)))

    try:
        storage_ok, storage_detail = _storage_check(collect_storage_report())
        checks.append(SelfTestCheck("Disk", storage_ok, _compact(storage_detail)))
    except Exception as exc:
        checks.append(SelfTestCheck("Disk", False, _compact(exc)))

    try:
        network_ok, network_detail = _network_check()
        checks.append(SelfTestCheck("Internet/Telegram", network_ok, _compact(network_detail)))
    except Exception as exc:
        checks.append(SelfTestCheck("Internet/Telegram", False, _compact(exc)))

    report = SelfTestReport(tuple(checks), time.monotonic() - started)
    if report.ok:
        log_info("[SELFTEST] All checks passed; no print job was submitted.")
    else:
        failed = "; ".join(f"{check.name}: {check.detail}" for check in report.checks if not check.ok)
        log_error(f"[SELFTEST] Failed; no print job was submitted. {_compact(failed, 1500)}")
    return report
