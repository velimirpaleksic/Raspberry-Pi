import hashlib
import re
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from project.core import config
from project.utils.docs.docx_replace_placeholders import replace_dynamic_text, value_fits_placeholder


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = PROJECT_ROOT / "project" / "docs" / "template.docx"
BACKUP = PROJECT_ROOT / "project" / "docs" / "template.2025-2026.original.backup.docx"
SECRETARY_BACKUP = PROJECT_ROOT / "project" / "docs" / "template.before-secretary-suzana-jokic.backup.docx"
ORIGINAL_SHA256 = "F2A99E1FEAF7827D81E226151470495135789412F442B9FA63B6C4B546C71E2B"
PRE_SECRETARY_SHA256 = "61FD522093EF16ADBDDA5F7A035C2B03A4A440B0E8B5065F243B29826B410C32"
OLD_SECRETARY_NAME = "Драгана Стојановић"
SECRETARY_NAME = "Сузана Јокић"
STUDENT_CAPTION = "                 (име и презиме ученика-це)                              (име  родитеља)"


def sample_placeholders():
    return {
        "{{DANASNJI_DATUM}}": "01.09.2026",
        "{{IME}}": "АЛЕКСАНДАР МАКСИМИЛИЈАН ПЕТРОВИЋ",
        "{{IME_PREZIME}}": "АЛЕКСАНДАР МАКСИМИЛИЈАН ПЕТРОВИЋ",
        "{{IME_UCENIKA}}": "АЛЕКСАНДАР",
        "{{PREZIME}}": "МАКСИМИЛИЈАН ПЕТРОВИЋ",
        "{{RODITELJ}}": "АЛЕКСАНДРИЈА",
        "{{DATUM_RODJENJA}}": "31.12.2008",
        "{{MJESTO}}": "ДОЊЕ НОВО СЕЛО",
        "{{OPSTINA}}": "ИСТОЧНА ИЛИЏА",
        "{{RAZRED}}": "IV-1",
        "{{STRUKA}}": max(config.STRUKE, key=len).upper(),
        "{{RAZLOG}}": max(config.RAZLOZI, key=len).upper(),
    }


class TemplateIntegrityTests(unittest.TestCase):
    def test_backup_hash_and_surgical_ooxml_diff(self):
        self.assertEqual(hashlib.sha256(BACKUP.read_bytes()).hexdigest().upper(), ORIGINAL_SHA256)
        with ZipFile(BACKUP) as old_zip, ZipFile(TEMPLATE) as new_zip:
            self.assertIsNone(old_zip.testzip())
            self.assertIsNone(new_zip.testzip())
            self.assertEqual(old_zip.namelist(), new_zip.namelist())
            changed = []
            for name in old_zip.namelist():
                if old_zip.read(name) != new_zip.read(name):
                    changed.append(name)
            self.assertEqual(changed, ["word/document.xml"])
            old_xml = old_zip.read("word/document.xml")
            new_xml = new_zip.read("word/document.xml")
            self.assertEqual(old_xml.count(b"2025/2026"), 1)
            self.assertEqual(new_xml.count(b"2025/2026"), 0)
            self.assertEqual(new_xml.count(b"2026/2027"), 1)
            self.assertEqual(sum(old_zip.read(name).count(b"2025/2026") for name in old_zip.namelist()), 1)
            self.assertEqual(sum(new_zip.read(name).count(b"2025/2026") for name in new_zip.namelist()), 0)
            self.assertEqual(sum(new_zip.read(name).count(b"2026/2027") for name in new_zip.namelist()), 1)

    def test_secretary_name_change_is_an_exact_ooxml_replacement(self):
        self.assertEqual(
            hashlib.sha256(SECRETARY_BACKUP.read_bytes()).hexdigest().upper(),
            PRE_SECRETARY_SHA256,
        )
        with ZipFile(SECRETARY_BACKUP) as old_zip, ZipFile(TEMPLATE) as new_zip:
            self.assertIsNone(old_zip.testzip())
            self.assertIsNone(new_zip.testzip())
            self.assertEqual(old_zip.namelist(), new_zip.namelist())
            changed = [
                name
                for name in old_zip.namelist()
                if old_zip.read(name) != new_zip.read(name)
            ]
            self.assertEqual(changed, ["word/document.xml"])
            old_xml = old_zip.read("word/document.xml")
            new_xml = new_zip.read("word/document.xml")
            old_name = OLD_SECRETARY_NAME.encode("utf-8")
            new_name = SECRETARY_NAME.encode("utf-8")
            self.assertEqual(old_xml.count(old_name), 1)
            self.assertEqual(old_xml.count(new_name), 0)
            self.assertEqual(new_xml, old_xml.replace(old_name, new_name, 1))
            self.assertEqual(new_xml.count(old_name), 0)
            self.assertEqual(new_xml.count(new_name), 1)

    def test_template_structure_and_static_student_caption_are_preserved(self):
        old_doc = Document(BACKUP)
        new_doc = Document(TEMPLATE)
        self.assertEqual(len(old_doc.sections), len(new_doc.sections))
        self.assertEqual(len(old_doc.paragraphs), len(new_doc.paragraphs))
        self.assertEqual(new_doc.paragraphs[7].text, STUDENT_CAPTION)
        changed = [
            i
            for i, (old_para, new_para) in enumerate(zip(old_doc.paragraphs, new_doc.paragraphs))
            if old_para.text != new_para.text
        ]
        self.assertEqual(changed, [13, 28])
        self.assertEqual(new_doc.paragraphs[13].text.count("2026/2027"), 1)
        self.assertIn(f"{SECRETARY_NAME}, дипл. правник", new_doc.paragraphs[28].text)
        self.assertNotIn(OLD_SECRETARY_NAME, new_doc.paragraphs[28].text)


class GeneratedDocumentTests(unittest.TestCase):
    def test_long_supported_values_generate_valid_docx_without_placeholders(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "output.docx"
            replace_dynamic_text(str(TEMPLATE), str(output), sample_placeholders())
            with ZipFile(output) as archive:
                self.assertIsNone(archive.testzip())
                unresolved = []
                for name in archive.namelist():
                    if name.endswith(".xml") and re.search(rb"\{\{[^{}]+\}\}", archive.read(name)):
                        unresolved.append(name)
                self.assertEqual(unresolved, [])

            generated = Document(output)
            self.assertEqual(len(generated.paragraphs), 31)
            self.assertEqual(generated.paragraphs[7].text, STUDENT_CAPTION)
            self.assertIn("2026/2027", generated.paragraphs[13].text)
            self.assertIn("АЛЕКСАНДАР\u00a0МАКСИМИЛИЈАН\u00a0ПЕТРОВИЋ", generated.paragraphs[5].text)
            self.assertIn(f"{SECRETARY_NAME}, дипл. правник", generated.paragraphs[28].text)

    def test_every_configured_reason_and_longest_profession_fit(self):
        self.assertTrue(all(value_fits_placeholder("{{RAZLOG}}", reason) for reason in config.RAZLOZI))
        self.assertTrue(value_fits_placeholder("{{STRUKA}}", max(config.STRUKE, key=len)))

    def test_unreadably_long_value_is_rejected(self):
        impossible = "ПРЕДУГАЧАК " * 30
        self.assertFalse(value_fits_placeholder("{{IME}}", impossible))
        values = sample_placeholders()
        values["{{IME}}"] = impossible
        with tempfile.TemporaryDirectory() as temp_dir:
            with self.assertRaises(ValueError):
                replace_dynamic_text(str(TEMPLATE), str(Path(temp_dir) / "invalid.docx"), values)


if __name__ == "__main__":
    unittest.main()
